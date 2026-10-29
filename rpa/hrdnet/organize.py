"""2 정리 — inbox 의 비용신청대상 엑셀 + 업체별 증빙 PDF → outbox manifest.json.

증빙 PDF는 outbox로 복사하지 않는다. inbox/증빙 아래의 업체별 원본 경로를 manifest에 직접 기록한다.
"""
from __future__ import annotations

import json
import re
import shutil
from datetime import datetime
from pathlib import Path

from . import config


def _read_rows(path: Path) -> list[dict]:
    if path.suffix.lower() == ".csv":
        import csv
        with open(path, encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f))
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        it = wb.active.iter_rows(values_only=True)
        head = [str(h).strip() if h is not None else "" for h in next(it, [])]
        return [dict(zip(head, r)) for r in it if any(v not in (None, "") for v in r)]
    finally:
        wb.close()  # read_only 는 파일을 잡고 있어서 닫아야 _archive 로 옮길 수 있다


def _norm(row: dict) -> dict:
    """엑셀 열 이름 → 코드 이름 (config.COLUMNS)."""
    out = {}
    for key, names in config.COLUMNS.items():
        for n in names:
            if n in row and row[n] not in (None, ""):
                v = row[n]
                if isinstance(v, datetime):
                    v = v.strftime("%Y-%m-%d")
                elif isinstance(v, float) and v.is_integer():
                    v = int(v)
                out[key] = str(v).strip() if key not in ("훈련인원", "신청금액") else v
                break
    return out


def _company_key(value: str) -> str:
    """파일명과 대시보드의 법인 표기 차이를 무시한 사업장 비교 키."""
    value = str(value or "")
    for token in ("주식회사", "유한회사", "(주)", "㈜", "(유한)", "(유)"):
        value = value.replace(token, "")
    return re.sub(r"[^0-9A-Za-z가-힣]", "", value).lower()


def _find_evidence(kind: str, rnd: str, course: str, given: str, pdfs: list[Path], base: Path,
                   companies: list[str]) -> Path | None:
    if given:
        hit = [p for p in pdfs if p.name == given or p.stem == Path(given).stem]
        if hit:
            return hit[0]
    num = re.compile(rf"(?<!\d){re.escape(str(rnd))}(?!\d)")
    words = config.EVIDENCE[kind]
    cands = [p for p in pdfs
             if num.search(str(p.relative_to(base))) and any(w in str(p.relative_to(base)) for w in words)]
    cands.sort(key=lambda p: (course not in str(p), -p.stat().st_mtime))  # 과정명이 들어간 것, 최근 것 먼저
    if len(cands) <= 1:
        return cands[0] if cands else None

    # 회차 폴더에 업체별 증빙이 여러 개 있으면 절대로 임의의 첫 파일을 고르지 않는다.
    # 현재 manifest가 사업장 1곳만 대상으로 할 때만 그 사업장명이 들어간 파일을 선택한다.
    company_keys = [k for k in (_company_key(name) for name in companies) if k]
    if len(company_keys) == 1:
        matched = [p for p in cands if company_keys[0] in _company_key(p.stem)]
        if len(matched) == 1:
            return matched[0]
    return None


def run(archive: bool = True) -> dict:
    d = config.dirs()
    sheets = [p for p in d["inbox"].iterdir() if p.is_file() and p.suffix.lower() in (".xlsx", ".csv")
              and not p.name.startswith("~$")]
    pdfs = [p for p in d["inbox"].rglob("*.pdf")]
    rounds: dict[tuple, dict] = {}
    problems: list[str] = []

    for sh in sheets:
        for i, raw in enumerate(_read_rows(sh), start=2):
            r = _norm(raw)
            miss = [k for k in config.REQUIRED if not r.get(k)]
            if miss:
                problems.append(f"{sh.name} {i}행: {', '.join(miss)} 없음 — 건너뜀")
                continue
            key = (r["과정명"], str(r["회차"]))
            m = rounds.setdefault(key, {
                "과정명": r["과정명"], "회차": str(r["회차"]),
                "검색시작": r.get("검색시작", ""), "검색종료": r.get("검색종료", ""),
                "사업장": [], "첨부": {}, "신청금액": 0, "훈련인원": 0,
                "원본": sh.name, "_given": {k: r.get(k, "") for k in config.EVIDENCE}})
            m["사업장"].append({"사업장명": r.get("사업장명", ""), "사업장관리번호": r.get("사업장관리번호", ""),
                               "훈련인원": r.get("훈련인원", 0), "신청금액": r.get("신청금액", 0)})
            m["신청금액"] += float(r.get("신청금액") or 0)
            m["훈련인원"] += int(float(r.get("훈련인원") or 0))

    made = []
    for (course, rnd), m in rounds.items():
        folder = d["outbox"] / f"{course}_{rnd}회차"
        folder.mkdir(parents=True, exist_ok=True)
        given = m.pop("_given")
        missing_evidence = []
        for kind in config.EVIDENCE:
            items = []
            for business in m["사업장"]:
                company = business.get("사업장명", "")
                row_given = given.get(kind, "") if len(m["사업장"]) == 1 else ""
                src = _find_evidence(kind, rnd, course, row_given, pdfs, d["inbox"], [company])
                if src:
                    items.append({"사업장명": company, "증빙종류": kind, "파일": str(src)})
                else:
                    missing_evidence.append(f"{company or '(사업장명 없음)'} - {kind}")
            m[f"{kind}목록"] = items
            m["첨부"][kind] = items[0]["파일"] if len(items) == 1 else ""
        if missing_evidence:
            problems.append(f"{course} {rnd}회차: 필수 증빙 없음 — {', '.join(missing_evidence)}")
        m["상태"] = "증빙누락" if missing_evidence else "준비완료"
        m["정리일시"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        (folder / "manifest.json").write_text(json.dumps(m, ensure_ascii=False, indent=1), encoding="utf-8")
        evidence_count = len(m["전자계산서목록"]) + len(m["수납확인서목록"])
        made.append(f"{folder.name}  사업장 {len(m['사업장'])}곳  첨부 {evidence_count}개  {m['상태']}")

    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    report = d["outbox"] / f"_검증리포트_{stamp}.txt"
    report.write_text("\n".join(["[정리된 회차]", *made, "", "[확인 필요]", *(problems or ["없음"])]),
                      encoding="utf-8")
    if archive and rounds:
        arc = d["_archive"] / stamp
        arc.mkdir(parents=True, exist_ok=True)
        for sh in sheets:
            shutil.move(str(sh), arc / sh.name)
    print(report.read_text(encoding="utf-8"))
    print(f"\n→ {d['outbox']}")
    return {"rounds": len(rounds), "problems": problems}
